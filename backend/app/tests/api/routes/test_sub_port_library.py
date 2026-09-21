"""Tests for sub-port-library API."""

import hashlib
import uuid
import zipfile
from collections.abc import Generator
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from PIL import Image
from sqlmodel import Session, delete, select

from app.core.config import settings
from app.core.db import engine
from app.models import ExportGroup, ExportGroupField, FileAttachment, SubPortRecord

TRACKED_GROUP_IDS: list[uuid.UUID] = []


@pytest.fixture(scope="module", autouse=True)
def _cleanup_sub_port_data() -> Generator[None, None, None]:
    yield
    with Session(engine) as session:
        session.execute(delete(SubPortRecord))
        for group_id in TRACKED_GROUP_IDS:
            group = session.get(ExportGroup, group_id)
            if group:
                session.delete(group)
        session.commit()


def _create_group(fields: list[tuple[str, str]]) -> dict:
    marker = uuid.uuid4().hex[:8]
    with Session(engine) as session:
        group = ExportGroup(name=f"子端口测试字段组{marker}")
        for sort_order, (code, label) in enumerate(fields):
            group.fields.append(
                ExportGroupField(
                    field_name=code, field_label=label, sort_order=sort_order
                )
            )
        session.add(group)
        session.commit()
        session.refresh(group)
        TRACKED_GROUP_IDS.append(group.id)
        return {"id": str(group.id), "name": group.name}


def _build_excel(headers: list[str], rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for col_idx, header in enumerate(headers, 1):
        ws.cell(row=1, column=col_idx, value=header)
    for row_idx, row in enumerate(rows, 2):
        for col_idx, value in enumerate(row, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def _import_file(
    client: TestClient,
    headers: dict[str, str],
    content: bytes,
    group_id: str,
    path: str = "/import",
):
    return client.post(
        f"{settings.API_V1_STR}/sub-port-library{path}",
        headers=headers,
        files={
            "file": (
                "data.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"group_id": group_id},
    )


def _template_headers(
    client: TestClient, headers: dict[str, str], group_id: str
) -> dict:
    r = client.get(
        f"{settings.API_V1_STR}/sub-port-library/template",
        headers=headers,
        params={"group_id": group_id},
    )
    assert r.status_code == 200, r.text
    wb = load_workbook(BytesIO(r.content))
    ws = wb.active
    return {
        "headers": [c.value for c in ws[1]],
        "sheets": wb.sheetnames,
        "data_rows": ws.max_row - 1,
    }


def test_download_template(
    client: TestClient, superuser_token_headers: dict[str, str]
) -> None:
    group = _create_group(
        [
            ("main_port_number", "主端口号"),
            ("operation_type", "操作类型"),
            ("sms_signature", "短信签名"),
            ("port_full_number", "短信子端口号"),
            ("enterprise_name", "企业名称"),
        ]
    )
    result = _template_headers(client, superuser_token_headers, group["id"])
    assert result["headers"] == [
        "状态",
        "操作类型",
        "子端口号",
        "主端口号",
        "短信子端口号",
        "短信签名",
        "企业名称",
        "是否四类",
        "其他举证图片",
        "子端口失败原因",
    ]
    assert "填写说明" in result["sheets"]
    # 模板不含示例数据行
    assert result["data_rows"] == 0


def test_import_creates_records(
    client: TestClient, superuser_token_headers: dict[str, str]
) -> None:
    group = _create_group([("sms_signature", "短信签名")])
    marker = uuid.uuid4().hex[:8]
    content = _build_excel(
        ["主端口号", "子端口号", "状态", "短信签名"],
        [[f"1069{marker}", f"8001{marker}", "", f"签名{marker}"]],
    )
    r = _import_file(client, superuser_token_headers, content, group["id"])
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["success_count"] == 1
    assert body["error_count"] == 0

    r = client.get(
        f"{settings.API_V1_STR}/sub-port-library",
        headers=superuser_token_headers,
        params={"keyword": marker},
    )
    assert r.status_code == 200
    data = r.json()["data"]
    assert len(data) == 1
    assert data[0]["status"] == "在线"
    assert data[0]["field_values"]["sms_signature"] == f"签名{marker}"
    assert "is_four_category" in data[0]["field_values"]
    assert "other_proof" in data[0]["field_values"]
    assert "sub_port_failure_reason" in data[0]["field_values"]


def test_import_duplicate_existing_is_rejected(
    client: TestClient, superuser_token_headers: dict[str, str]
) -> None:
    group = _create_group([("sms_signature", "短信签名")])
    marker = uuid.uuid4().hex[:8]
    main = f"1069{marker}"
    sub = f"9001{marker}"
    first = _build_excel(
        ["主端口号", "子端口号", "状态", "短信签名"],
        [[main, sub, "在线", f"旧签名{marker}"]],
    )
    assert (
        _import_file(client, superuser_token_headers, first, group["id"]).json()[
            "success_count"
        ]
        == 1
    )

    second = _build_excel(
        ["主端口号", "子端口号", "状态", "短信签名"],
        [[main, sub, "整改", f"新签名{marker}"]],
    )
    r = _import_file(client, superuser_token_headers, second, group["id"])
    assert r.status_code == 200
    body = r.json()
    assert body["success_count"] == 0
    assert body["error_count"] == 1
    assert body["errors"][0]["row"] == 2
    assert body["errors"][0]["reason"] == "子端口号已存在，导入不会覆盖已有记录"

    r = client.get(
        f"{settings.API_V1_STR}/sub-port-library",
        headers=superuser_token_headers,
        params={"keyword": marker},
    )
    data = r.json()["data"]
    assert len(data) == 1
    assert data[0]["status"] == "在线"
    assert data[0]["field_values"]["sms_signature"] == f"旧签名{marker}"


def test_import_validation_errors(
    client: TestClient, superuser_token_headers: dict[str, str]
) -> None:
    group = _create_group([])
    marker = uuid.uuid4().hex[:8]
    content = _build_excel(
        ["主端口号", "子端口号", "状态"],
        [
            ["", f"8001{marker}", "在线"],
            [f"1069{marker}", "", "在线"],
            [f"1069{marker}", f"8002{marker}", "停机"],
            [f"1069{marker}", f"8003{marker}", ""],
        ],
    )
    r = _import_file(client, superuser_token_headers, content, group["id"])
    assert r.status_code == 200
    body = r.json()
    assert body["error_count"] == 3
    assert body["success_count"] == 1
    assert {e["row"] for e in body["errors"]} == {2, 3, 4}

    # 错误行不入库
    r = client.get(
        f"{settings.API_V1_STR}/sub-port-library",
        headers=superuser_token_headers,
        params={"keyword": marker},
    )
    assert len(r.json()["data"]) == 1


def test_import_duplicate_in_file(
    client: TestClient, superuser_token_headers: dict[str, str]
) -> None:
    group = _create_group([])
    marker = uuid.uuid4().hex[:8]
    main = f"1069{marker}"
    content = _build_excel(
        ["主端口号", "子端口号", "状态"],
        [
            [main, f"8001{marker}", "在线"],
            [main, f"8001{marker}", "下线"],
        ],
    )
    r = _import_file(client, superuser_token_headers, content, group["id"])
    assert r.status_code == 200
    body = r.json()
    assert body["success_count"] == 1
    assert body["error_count"] == 1
    assert body["errors"][0]["row"] == 3


def test_create_duplicate_400(
    client: TestClient, superuser_token_headers: dict[str, str]
) -> None:
    marker = uuid.uuid4().hex[:8]
    payload = {
        "main_port_number": f"1069{marker}",
        "sub_port_number": f"8001{marker}",
    }
    r = client.post(
        f"{settings.API_V1_STR}/sub-port-library",
        headers=superuser_token_headers,
        json=payload,
    )
    assert r.status_code == 200, r.text
    record_id = r.json()["id"]

    r = client.post(
        f"{settings.API_V1_STR}/sub-port-library",
        headers=superuser_token_headers,
        json=payload,
    )
    assert r.status_code == 400

    # 清理
    r = client.delete(
        f"{settings.API_V1_STR}/sub-port-library/{record_id}",
        headers=superuser_token_headers,
    )
    assert r.status_code == 200


def test_update_duplicate_400(
    client: TestClient, superuser_token_headers: dict[str, str]
) -> None:
    marker = uuid.uuid4().hex[:8]
    ids = []
    for sub in (f"8001{marker}", f"8002{marker}"):
        r = client.post(
            f"{settings.API_V1_STR}/sub-port-library",
            headers=superuser_token_headers,
            json={"main_port_number": f"1069{marker}", "sub_port_number": sub},
        )
        assert r.status_code == 200, r.text
        ids.append(r.json()["id"])

    # 把第二条改成与第一条相同 → 400
    r = client.patch(
        f"{settings.API_V1_STR}/sub-port-library/{ids[1]}",
        headers=superuser_token_headers,
        json={"sub_port_number": f"8001{marker}"},
    )
    assert r.status_code == 400

    # 非法状态 → 400
    r = client.patch(
        f"{settings.API_V1_STR}/sub-port-library/{ids[1]}",
        headers=superuser_token_headers,
        json={"status": "停机"},
    )
    assert r.status_code == 400

    for record_id in ids:
        client.delete(
            f"{settings.API_V1_STR}/sub-port-library/{record_id}",
            headers=superuser_token_headers,
        )


def test_delete_list_parse_and_delete(
    client: TestClient, superuser_token_headers: dict[str, str]
) -> None:
    group = _create_group([])
    marker = uuid.uuid4().hex[:8]
    rows = [
        [f"1069{marker}", f"8001{marker}", "在线"],
        [f"1069{marker}", f"8002{marker}", "在线"],
    ]
    content = _build_excel(["主端口号", "子端口号", "状态"], rows)
    assert (
        _import_file(client, superuser_token_headers, content, group["id"]).json()[
            "success_count"
        ]
        == 2
    )

    delete_content = _build_excel(
        ["主端口号", "子端口号"],
        [
            [f"1069{marker}", f"8001{marker}"],
            [f"1069{marker}", f"9999{marker}"],
        ],
    )
    r = _import_file(
        client,
        superuser_token_headers,
        delete_content,
        group["id"],
        "/import/parse-delete",
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["matched_count"] == 1
    assert body["total"] == 2
    assert body["unmatched"][0]["sub_port_number"] == f"9999{marker}"

    r = _import_file(
        client, superuser_token_headers, delete_content, group["id"], "/import/delete"
    )
    assert r.status_code == 200
    body = r.json()
    assert body["deleted_count"] == 1
    assert len(body["unmatched"]) == 1

    r = client.get(
        f"{settings.API_V1_STR}/sub-port-library",
        headers=superuser_token_headers,
        params={"main_port_number": f"1069{marker}"},
    )
    remaining = r.json()["data"]
    assert len(remaining) == 1
    assert remaining[0]["sub_port_number"] == f"8002{marker}"


def test_list_filters(
    client: TestClient, superuser_token_headers: dict[str, str]
) -> None:
    group = _create_group([("sms_signature", "短信签名")])
    marker = uuid.uuid4().hex[:8]
    main = f"1069{marker}"
    rows = [
        [main, f"8001{marker}", "在线", f"签名A{marker}"],
        [main, f"8002{marker}", "下线", f"签名B{marker}"],
        [f"1070{marker}", f"8003{marker}", "整改", f"签名C{marker}"],
    ]
    content = _build_excel(["主端口号", "子端口号", "状态", "短信签名"], rows)
    assert (
        _import_file(client, superuser_token_headers, content, group["id"]).json()[
            "success_count"
        ]
        == 3
    )

    # keyword 模糊
    r = client.get(
        f"{settings.API_V1_STR}/sub-port-library",
        headers=superuser_token_headers,
        params={"keyword": f"8001{marker}"},
    )
    assert len(r.json()["data"]) == 1

    # keyword 命中字段组动态字段
    r = client.get(
        f"{settings.API_V1_STR}/sub-port-library",
        headers=superuser_token_headers,
        params={"keyword": f"签名B{marker}"},
    )
    data = r.json()["data"]
    assert len(data) == 1
    assert data[0]["sub_port_number"] == f"8002{marker}"

    # status 精确
    r = client.get(
        f"{settings.API_V1_STR}/sub-port-library",
        headers=superuser_token_headers,
        params={"status": "下线", "main_port_number": main},
    )
    data = r.json()["data"]
    assert len(data) == 1
    assert data[0]["sub_port_number"] == f"8002{marker}"

    # main_port_number 精确
    r = client.get(
        f"{settings.API_V1_STR}/sub-port-library",
        headers=superuser_token_headers,
        params={"main_port_number": f"1070{marker}"},
    )
    assert len(r.json()["data"]) == 1

    # 分页
    r = client.get(
        f"{settings.API_V1_STR}/sub-port-library",
        headers=superuser_token_headers,
        params={"main_port_number": main, "page": 1, "page_size": 1},
    )
    body = r.json()
    assert body["total"] == 2
    assert len(body["data"]) == 1


def test_export_selected_fields(
    client: TestClient, superuser_token_headers: dict[str, str]
) -> None:
    group = _create_group(
        [("sms_signature", "短信签名"), ("enterprise_name", "企业名称")]
    )
    marker = uuid.uuid4().hex[:8]
    main = f"1069{marker}"
    content = _build_excel(
        ["主端口号", "子端口号", "状态", "短信签名", "企业名称"],
        [
            [main, f"8001{marker}", "在线", f"签名A{marker}", f"企业A{marker}"],
            [main, f"8002{marker}", "在线", f"签名B{marker}", f"企业B{marker}"],
        ],
    )
    assert (
        _import_file(client, superuser_token_headers, content, group["id"]).json()[
            "success_count"
        ]
        == 2
    )

    r = client.get(
        f"{settings.API_V1_STR}/sub-port-library/export",
        headers=superuser_token_headers,
        params={
            "group_id": group["id"],
            "main_port_number": main,
            "field_names": "sms_signature",
        },
    )
    assert r.status_code == 200, r.text
    wb = load_workbook(BytesIO(r.content))
    ws = wb.active
    assert [c.value for c in ws[1]] == ["状态", "子端口号", "主端口号", "短信签名"]
    assert ws.max_row == 3
    assert ws.cell(row=2, column=4).value in {f"签名A{marker}", f"签名B{marker}"}


def test_export_embeds_attachment_instead_of_dispimg_formula(
    monkeypatch, client: TestClient, superuser_token_headers: dict[str, str]
) -> None:
    group = _create_group([])
    marker = uuid.uuid4().hex[:8]
    main = f"1069{marker}"
    sub = f"8001{marker}"
    r = client.post(
        f"{settings.API_V1_STR}/sub-port-library",
        headers=superuser_token_headers,
        json={
            "main_port_number": main,
            "sub_port_number": sub,
            "field_values": {"other_proof": '=DISPIMG("image-id")'},
        },
    )
    assert r.status_code == 200, r.text
    record_id = uuid.UUID(r.json()["id"])

    image_output = BytesIO()
    Image.new("RGB", (40, 20), "red").save(image_output, format="PNG")
    image_bytes = image_output.getvalue()
    stored_path = f"sub_port_record_images/{marker}.png"
    with Session(engine) as session:
        session.add(
            FileAttachment(
                original_name="image_row2_col8.png",
                stored_path=stored_path,
                file_size=len(image_bytes),
                mime_type="image/png",
                md5_hash=hashlib.md5(image_bytes).hexdigest(),
                entity_type="sub_port_record",
                entity_id=record_id,
                field_name="其他举证图片",
            )
        )
        session.commit()

    class FakeStorage:
        def download(self, key: str) -> bytes:
            assert key == stored_path
            return image_bytes

    monkeypatch.setattr(
        "app.api.routes.sub_port_library.get_storage", lambda: FakeStorage()
    )
    r = client.get(
        f"{settings.API_V1_STR}/sub-port-library/export",
        headers=superuser_token_headers,
        params={
            "group_id": group["id"],
            "main_port_number": main,
            "field_names": "other_proof",
        },
    )
    assert r.status_code == 200, r.text
    wb = load_workbook(BytesIO(r.content))
    ws = wb.active
    assert [cell.value for cell in ws[1]] == [
        "状态",
        "子端口号",
        "主端口号",
        "其他举证图片",
    ]
    assert ws.cell(row=2, column=4).value is None
    assert len(ws._images) == 1
    with zipfile.ZipFile(BytesIO(r.content)) as archive:
        names = set(archive.namelist())
        assert "xl/drawings/drawing1.xml" in names
        assert any(name.startswith("xl/media/image") for name in names)
        assert b"DISPIMG" not in archive.read("xl/worksheets/sheet1.xml")

    with Session(engine) as session:
        attachment = session.exec(
            select(FileAttachment).where(FileAttachment.entity_id == record_id)
        ).one()
        session.delete(attachment)
        session.commit()


def test_manual_crud_and_batch_delete(
    client: TestClient, superuser_token_headers: dict[str, str]
) -> None:
    marker = uuid.uuid4().hex[:8]
    main = f"1069{marker}"

    r = client.post(
        f"{settings.API_V1_STR}/sub-port-library",
        headers=superuser_token_headers,
        json={
            "main_port_number": main,
            "sub_port_number": f"8001{marker}",
            "status": "整改",
            "field_values": {"sms_signature": f"签名{marker}"},
        },
    )
    assert r.status_code == 200, r.text
    created = r.json()
    assert created["status"] == "整改"
    assert created["field_values"]["sms_signature"] == f"签名{marker}"
    record_id = created["id"]

    r = client.patch(
        f"{settings.API_V1_STR}/sub-port-library/{record_id}",
        headers=superuser_token_headers,
        json={"status": "下线", "field_values": {"sms_signature": ""}},
    )
    assert r.status_code == 200, r.text
    updated = r.json()
    assert updated["status"] == "下线"
    assert updated["field_values"]["sms_signature"] == ""

    r = client.post(
        f"{settings.API_V1_STR}/sub-port-library",
        headers=superuser_token_headers,
        json={"main_port_number": main, "sub_port_number": f"8002{marker}"},
    )
    assert r.status_code == 200, r.text
    second_id = r.json()["id"]

    r = client.post(
        f"{settings.API_V1_STR}/sub-port-library/batch-delete",
        headers=superuser_token_headers,
        json={"ids": [record_id, second_id]},
    )
    assert r.status_code == 200
    assert r.json()["deleted_count"] == 2

    r = client.get(
        f"{settings.API_V1_STR}/sub-port-library",
        headers=superuser_token_headers,
        params={"keyword": marker},
    )
    assert len(r.json()["data"]) == 0
