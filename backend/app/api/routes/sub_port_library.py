"""子端口库 API 路由。"""

import io
import uuid
from typing import Any
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook

from app.api.deps import SessionDep, get_current_active_superuser
from app.crud.export_group import get_export_group
from app.crud.sub_port_record import (
    create_sub_port_record,
    delete_sub_port_record,
    delete_sub_port_records,
    get_by_main_and_sub,
    get_sub_port_record,
    list_sub_port_records,
    update_sub_port_record,
)
from app.models import (
    SUB_PORT_STATUSES,
    ExportGroup,
    Message,
    SubPortBatchDelete,
    SubPortRecord,
    SubPortRecordCreate,
    SubPortRecordPublic,
    SubPortRecordsPublic,
    SubPortRecordUpdate,
)

router = APIRouter(
    prefix="/sub-port-library",
    tags=["sub-port-library"],
    dependencies=[Depends(get_current_active_superuser)],
)

FIXED_HEADERS = ("主端口号", "子端口号", "状态")
DELETE_MAIN_ALIASES = {"主端口号", "主端口"}
DELETE_SUB_ALIASES = {"子端口号", "子端口"}


def _cell_text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _load_group(session: SessionDep, group_id: uuid.UUID) -> ExportGroup:
    group = get_export_group(session=session, id=group_id)
    if not group:
        raise HTTPException(status_code=404, detail="字段组不存在")
    return group


def _validate_status(status: str) -> str:
    if status not in SUB_PORT_STATUSES:
        raise HTTPException(
            status_code=400,
            detail=f"状态必须是 {'、'.join(SUB_PORT_STATUSES)} 之一",
        )
    return status


def _check_unique(
    session: SessionDep,
    main_port_number: str,
    sub_port_number: str,
    exclude_id: uuid.UUID | None = None,
) -> None:
    existing = get_by_main_and_sub(
        session=session,
        main_port_number=main_port_number,
        sub_port_number=sub_port_number,
    )
    if existing and existing.id != exclude_id:
        raise HTTPException(
            status_code=400,
            detail=(
                f"子端口号已存在：主端口号 {main_port_number} "
                f"下已存在子端口号 {sub_port_number}"
            ),
        )


def _sorted_fields(group: ExportGroup) -> list:
    return sorted(group.fields, key=lambda f: f.sort_order)


def _header_to_field(group: ExportGroup) -> dict[str, str]:
    """表头文本 -> 数据键（固定列 + 字段组字段 label）。"""
    mapping: dict[str, str] = {}
    for header in FIXED_HEADERS:
        mapping.setdefault(header, header)
    for field in _sorted_fields(group):
        mapping.setdefault(field.field_label, field.field_name)
    return mapping


def _parse_sub_port_excel(
    content: bytes, group: ExportGroup
) -> tuple[list[str], list[dict], list[dict], int]:
    """解析导入 Excel，返回 (headers, records, errors, total_data_rows)。"""
    try:
        wb = load_workbook(io.BytesIO(content))
    except Exception:
        raise HTTPException(status_code=400, detail="无法解析 Excel 文件，请检查文件格式")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 1:
        raise HTTPException(status_code=400, detail="文件为空，请导入有效的 Excel 文件")

    headers = [_cell_text(c) for c in rows[0]]
    header_to_field = _header_to_field(group)

    col_map: dict[str, int] = {}
    for col_idx, header in enumerate(headers):
        if header in header_to_field:
            col_map.setdefault(header_to_field[header], col_idx)

    for required in ("主端口号", "子端口号"):
        if required not in col_map:
            raise HTTPException(
                status_code=400,
                detail=f"缺少必要的表头列：{required}，请使用导入模板",
            )

    group_fields = _sorted_fields(group)
    records: list[dict] = []
    errors: list[dict] = []
    seen: set[tuple[str, str]] = set()
    total_data_rows = 0

    for row_idx, row in enumerate(rows[1:], start=2):
        values = [_cell_text(c) for c in row]
        if all(v == "" for v in values):
            continue
        total_data_rows += 1

        main_port_number = values[col_map["主端口号"]] if col_map["主端口号"] < len(values) else ""
        sub_port_number = values[col_map["子端口号"]] if col_map["子端口号"] < len(values) else ""
        status = values[col_map["状态"]] if col_map.get("状态") is not None and col_map["状态"] < len(values) else ""

        if not main_port_number:
            errors.append(
                {
                    "row": row_idx,
                    "field": "主端口号",
                    "value": "",
                    "reason": "主端口号为空",
                    "suggestion": "请填写主端口号",
                }
            )
            continue
        if not sub_port_number:
            errors.append(
                {
                    "row": row_idx,
                    "field": "子端口号",
                    "value": "",
                    "reason": "子端口号为空",
                    "suggestion": "请填写子端口号",
                }
            )
            continue
        if status and status not in SUB_PORT_STATUSES:
            errors.append(
                {
                    "row": row_idx,
                    "field": "状态",
                    "value": status,
                    "reason": f"状态必须是 {'、'.join(SUB_PORT_STATUSES)} 之一",
                    "suggestion": "请填写：在线、下线 或 整改",
                }
            )
            continue
        key = (main_port_number, sub_port_number)
        if key in seen:
            errors.append(
                {
                    "row": row_idx,
                    "field": "主端口号+子端口号",
                    "value": f"{main_port_number} / {sub_port_number}",
                    "reason": "文件内存在重复的主端口号+子端口号",
                    "suggestion": "同一组合仅保留一行，请删除重复行",
                }
            )
            continue
        if not status:
            status = "在线"
        seen.add(key)

        field_values: dict[str, str] = {}
        for field in group_fields:
            col_idx = col_map.get(field.field_name)
            value = values[col_idx] if col_idx is not None and col_idx < len(values) else ""
            field_values[field.field_name] = value

        records.append(
            {
                "main_port_number": main_port_number,
                "sub_port_number": sub_port_number,
                "status": status,
                "field_values": field_values,
            }
        )

    return headers, records, errors, total_data_rows


def _parse_delete_list(content: bytes) -> list[tuple[str, str]]:
    """解析删除清单，返回去重后的 (主端口号, 子端口号) 列表。"""
    try:
        wb = load_workbook(io.BytesIO(content))
    except Exception:
        raise HTTPException(status_code=400, detail="无法解析 Excel 文件，请检查文件格式")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 1:
        raise HTTPException(status_code=400, detail="文件为空")

    main_col: int | None = None
    sub_col: int | None = None
    data_start = 0
    header_row = [_cell_text(c) for c in rows[0]]
    for col_idx, header in enumerate(header_row[:4]):
        if header in DELETE_MAIN_ALIASES and main_col is None:
            main_col = col_idx
        elif header in DELETE_SUB_ALIASES and sub_col is None:
            sub_col = col_idx
    if main_col is not None or sub_col is not None:
        data_start = 1
    if main_col is None:
        main_col = 0
    if sub_col is None:
        sub_col = 1

    pairs: list[tuple[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for row in rows[data_start:]:
        values = [_cell_text(c) for c in row]
        if all(v == "" for v in values):
            continue
        main_port_number = values[main_col] if main_col < len(values) else ""
        sub_port_number = values[sub_col] if sub_col is not None and sub_col < len(values) else ""
        if not main_port_number or not sub_port_number:
            continue
        key = (main_port_number, sub_port_number)
        if key not in seen:
            seen.add(key)
            pairs.append(key)
    return pairs


def _target_detail(target: tuple[str, str]) -> dict[str, str]:
    return {
        "main_port_number": target[0],
        "sub_port_number": target[1],
    }


@router.get("/template")
def download_template(*, session: SessionDep, group_id: uuid.UUID) -> Any:
    group = _load_group(session, group_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "子端口数据"
    headers = list(FIXED_HEADERS) + [
        field.field_label for field in _sorted_fields(group)
    ]
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)

    instructions = wb.create_sheet("填写说明")
    instructions.cell(row=1, column=1, value="子端口数据导入填写说明")
    notes = [
        "1. 主端口号、子端口号为必填项，不能为空；",
        "2. 状态为单选：在线 / 下线 / 整改，留空默认为“在线”；",
        f"3. 其余列为当前字段组“{group.name}”的自定义字段，选填；",
        "4. 导入时按“主端口号+子端口号”匹配：已存在则覆盖更新，不存在则新增；",
        "5. 同一文件内不允许出现重复的“主端口号+子端口号”组合。",
    ]
    for row_idx, note in enumerate(notes, 2):
        instructions.cell(row=row_idx, column=1, value=note)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{quote('子端口数据导入模板.xlsx')}"
        },
    )


@router.post("/import/preview")
def preview_import(
    *, session: SessionDep, file: UploadFile = File(...), group_id: uuid.UUID = Form(...)
) -> Any:
    group = _load_group(session, group_id)
    if not file.filename or not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 或 .xls 文件")
    headers, records, errors, total_data_rows = _parse_sub_port_excel(
        file.file.read(), group
    )

    header_to_field = _header_to_field(group)
    recognized_headers = [h for h in headers if h and h in header_to_field]
    unrecognized = [
        h for h in headers if h and h not in header_to_field and h not in ("", "None")
    ]

    preview_rows = []
    for record in records[:5]:
        row_data: dict[str, str] = {}
        for header in recognized_headers:
            key = header_to_field[header]
            if key == "主端口号":
                row_data[header] = record["main_port_number"]
            elif key == "子端口号":
                row_data[header] = record["sub_port_number"]
            elif key == "状态":
                row_data[header] = record["status"]
            else:
                row_data[header] = record["field_values"].get(key, "")
        preview_rows.append(row_data)

    return {
        "headers": headers,
        "rows": preview_rows,
        "unrecognized_headers": unrecognized,
        "total_data_rows": total_data_rows,
    }


@router.post("/import")
def import_sub_ports(
    *, session: SessionDep, file: UploadFile = File(...), group_id: uuid.UUID = Form(...)
) -> Any:
    group = _load_group(session, group_id)
    if not file.filename or not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 或 .xls 文件")
    _, records, errors, total_data_rows = _parse_sub_port_excel(
        file.file.read(), group
    )

    for record in records:
        existing = get_by_main_and_sub(
            session=session,
            main_port_number=record["main_port_number"],
            sub_port_number=record["sub_port_number"],
        )
        if existing:
            existing.status = record["status"]
            existing.field_values = record["field_values"]
            session.add(existing)
        else:
            session.add(SubPortRecord(**record))
    session.commit()

    success_count = len(records)
    return {
        "total": total_data_rows,
        "success_count": success_count,
        "error_count": len(errors),
        "errors": errors,
        "message": f"导入完成：成功 {success_count} 条（新增或覆盖更新），失败 {len(errors)} 条",
    }


@router.post("/import/parse-delete")
def parse_delete_list_endpoint(*, session: SessionDep, file: UploadFile = File(...)) -> Any:
    if not file.filename or not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 或 .xls 文件")
    pairs = _parse_delete_list(file.file.read())

    matched: list[dict[str, str]] = []
    unmatched: list[dict[str, str]] = []
    for pair in pairs:
        existing = get_by_main_and_sub(
            session=session,
            main_port_number=pair[0],
            sub_port_number=pair[1],
        )
        (matched if existing else unmatched).append(_target_detail(pair))

    return {
        "matched_count": len(matched),
        "matched": matched,
        "unmatched": unmatched,
        "total": len(pairs),
    }


@router.post("/import/delete")
def delete_by_list_endpoint(*, session: SessionDep, file: UploadFile = File(...)) -> Any:
    if not file.filename or not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 或 .xls 文件")
    pairs = _parse_delete_list(file.file.read())

    deleted_count = 0
    unmatched: list[dict[str, str]] = []
    for pair in pairs:
        existing = get_by_main_and_sub(
            session=session,
            main_port_number=pair[0],
            sub_port_number=pair[1],
        )
        if existing:
            session.delete(existing)
            deleted_count += 1
        else:
            unmatched.append(_target_detail(pair))
    session.commit()

    return {
        "deleted_count": deleted_count,
        "unmatched": unmatched,
        "total": len(pairs),
    }


@router.get("", response_model=SubPortRecordsPublic)
@router.get("/", include_in_schema=False, response_model=SubPortRecordsPublic)
def read_sub_port_records(
    session: SessionDep,
    page: int = 1,
    page_size: int = 20,
    keyword: str | None = None,
    status: str | None = None,
    main_port_number: str | None = None,
) -> Any:
    records, count = list_sub_port_records(
        session=session,
        page=page,
        page_size=page_size,
        keyword=keyword,
        status=status,
        main_port_number=main_port_number,
    )
    return SubPortRecordsPublic(
        data=records, total=count, page=page, page_size=page_size
    )


@router.post("", response_model=SubPortRecordPublic)
@router.post("/", include_in_schema=False, response_model=SubPortRecordPublic)
def create_sub_port_record_endpoint(
    *, session: SessionDep, create: SubPortRecordCreate
) -> Any:
    _validate_status(create.status)
    _check_unique(session, create.main_port_number, create.sub_port_number)
    return create_sub_port_record(session=session, create=create)


@router.patch("/{id}", response_model=SubPortRecordPublic)
def update_sub_port_record_endpoint(
    *, session: SessionDep, id: uuid.UUID, update: SubPortRecordUpdate
) -> Any:
    db_obj = get_sub_port_record(session=session, id=id)
    if not db_obj:
        raise HTTPException(status_code=404, detail="子端口记录不存在")
    data = update.model_dump(exclude_unset=True)
    if "status" in data:
        _validate_status(data["status"])
    new_main = data.get("main_port_number", db_obj.main_port_number)
    new_sub = data.get("sub_port_number", db_obj.sub_port_number)
    _check_unique(session, new_main, new_sub, exclude_id=id)
    return update_sub_port_record(session=session, db_obj=db_obj, update=update)


@router.delete("/{id}")
def delete_sub_port_record_endpoint(
    *, session: SessionDep, id: uuid.UUID
) -> Message:
    db_obj = get_sub_port_record(session=session, id=id)
    if not db_obj:
        raise HTTPException(status_code=404, detail="子端口记录不存在")
    delete_sub_port_record(session=session, db_obj=db_obj)
    return Message(message="子端口记录删除成功")


@router.post("/batch-delete")
def batch_delete_sub_port_records_endpoint(
    *, session: SessionDep, body: SubPortBatchDelete
) -> Any:
    deleted_count = delete_sub_port_records(session=session, ids=body.ids)
    return {"deleted_count": deleted_count}
